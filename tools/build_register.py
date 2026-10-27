"""Build the job-card register and the unit reconciliation for one incoming batch.

    python3 tools/build_register.py <work>/docx incoming/2026-10-27/.../توزيع العاملين.xlsx \
        analysis/unit_map.csv analysis/2026-10-27_register.xlsx

Three sources are joined on analysis/unit_map.csv, which is maintained by hand:
the bylaws article of each unit, the card file that describes its jobs, and its
row in the staff-distribution sheet. Nothing is matched on names automatically,
because the three sources spell the same unit differently.
"""
import csv, re, sys, pathlib, collections, unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from parse_cards import cards

docx_root, staff_xlsx, map_csv, dst = map(pathlib.Path, sys.argv[1:5])


def card_key(f):
    """Article key of a card file: '16 .../3 بطاقات' -> 16.3, '20-1 بطاقات' -> 20."""
    nf = lambda x: unicodedata.normalize('NFKC', x)  # some names arrive in presentation forms
    stem, parent = nf(f.stem), nf(f.parent.name)
    if 'أمين السر' in stem: return '16.0'
    if stem.startswith('2020'): return '2020'
    n = re.match(r'(\d+)', stem).group(1)
    return f'16.{n}' if parent.startswith('16 ') else n


def level(s):
    s = s or ''
    if 'عليا' in s: return 'إدارة عليا'
    if 'وسط' in s: return 'إدارة وسطى'
    if re.search(r'^[اإ]د(ا|ر)', s): return 'إداري'
    if 'فني' in s: return 'فني'
    if re.search('خدم|عضل', s): return 'خدمي'
    return 'غير محدد'


def qual(s):
    for k, w in [('دراسات عليا', 'ماجستير|دكتوراه|دراسات عليا'), ('جامعي', 'جامع|هندس|إجازة|اجازة'),
                 ('كلية عسكرية / شرطة', 'عسكري|شرطة'), ('معهد', 'معهد'), ('ثانوي', 'ثانوي'),
                 ('إعدادي', 'اعدادي|إعدادي'), ('ابتدائي', 'ابتدائ|محو')]:
        if re.search(w, s or ''): return k
    return 'غير محدد'


def years(s):
    m = re.search(r'(\d+)', s or '')
    return int(m.group(1)) if m else None


# cards
rows = []
for f in sorted(docx_root.rglob('*.docx')):
    st = unicodedata.normalize('NFKC', f.stem)
    if 'بطاق' not in st and not st.startswith('2020'): continue
    for i, c in enumerate(cards(f), 1):
        rows.append({'article': card_key(f), 'file': f.stem, 'n': i, **c,
                     'level': level(c['cls']), 'qual_level': qual(c['qual']),
                     'exp_years': years(c['exp'])})

# staff sheet: row label -> 7 education counts
ws = openpyxl.load_workbook(staff_xlsx, data_only=True).worksheets[0]
EDU = ['هندسية', 'إنسانية', 'أساسية', 'طبية', 'معهد', 'ثانوية عامة', 'ثانوية مهنية']
staff = {r[0].strip(): list(r[1:8]) for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}

umap = list(csv.DictReader(open(map_csv, encoding='utf-8')))
by_art = collections.defaultdict(list)
for r in rows: by_art[r['article']].append(r)

wb = openpyxl.Workbook()
bold, head = Font(bold=True), PatternFill('solid', fgColor='DDE7F0')

def sheet(ws, header, data, widths):
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for c in ws[1]: c.font, c.fill = bold, head
    for d in data: ws.append(d)
    for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

u = wb.active; u.title = 'الوحدات'
udata = []
for m in umap:
    cs = by_art.get(m['card_key'], []) if m['card_key'] else []
    st = staff.get(m['staff_row']) if m['staff_row'] else None
    lv = collections.Counter(c['level'] for c in cs)
    udata.append([m['article'], m['unit_bylaws'], m['staff_row'], m['match'], len(cs),
                  len({c['title'] for c in cs}), lv['إدارة عليا'], lv['إدارة وسطى'],
                  sum(st[:4]) if st else None, sum(st[4:]) if st else None, sum(st) if st else None,
                  m['note']])
sheet(u, ['المادة', 'الوحدة في النظام الداخلي', 'السطر في جدول العاملين', 'المطابقة', 'عدد البطاقات',
          'مسميات مختلفة', 'إدارة عليا', 'إدارة وسطى', 'عاملون جامعيون', 'عاملون دون الجامعي',
          'مجموع العاملين', 'ملاحظة'], udata, [8, 42, 28, 10, 11, 12, 10, 10, 13, 15, 13, 70])

c = wb.create_sheet('البطاقات')
sheet(c, ['المادة', 'الملف', 'رقم', 'المديرية', 'الدائرة', 'الشعبة', 'الرئيس المباشر', 'المسمى',
          'التصنيف كما ورد', 'التصنيف', 'المؤهل كما ورد', 'المؤهل', 'الاختصاص', 'الخبرة كما وردت',
          'سنوات الخبرة', 'نوع الجهد', 'عدد المهام', 'المهام'],
      [[r['article'], r['file'], r['n'], r['dir'], r['dept'], r['sec'], r['boss'], r['title'],
        r['cls'], r['level'], r['qual'], r['qual_level'], r['spec'], r['exp'], r['exp_years'],
        r['effort'], len(r['tasks']), '\n'.join(r['tasks'])] for r in rows],
      [8, 22, 6, 26, 26, 22, 22, 30, 16, 12, 22, 12, 22, 16, 10, 14, 9, 90])
for row in c.iter_rows(min_row=2, min_col=18, max_col=18):
    row[0].alignment = Alignment(wrap_text=True, vertical='top')

s = wb.create_sheet('جدول العاملين')
sheet(s, ['السطر', 'علوم هندسية', 'علوم إنسانية', 'علوم أساسية', 'علوم طبية', 'معهد',
          'ثانوية عامة', 'ثانوية مهنية', 'المجموع', 'مطابق لوحدة في النظام؟'],
      [[k, *v, sum(v), next((m['match'] for m in umap if m['staff_row'] == k), 'no')]
       for k, v in staff.items()], [30, 11, 11, 11, 10, 8, 11, 11, 9, 20])

wb.save(dst)
print(f'{len(rows)} cards, {len(umap)} unit rows, {len(staff)} staff rows -> {dst}')
