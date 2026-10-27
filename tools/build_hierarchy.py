"""Derive the organisation structure from the job cards alone.

    python3 tools/build_hierarchy.py <work>/docx analysis/unit_map.csv analysis/<date>_hierarchy.xlsx

A card says where a job sits (المديرية / الدائرة / الشعبة / قسم-مكتب-فرع) and
who it reports to (الرئيس المباشر). The supervisor is a ROLE, not a name:
"رئيس الدائرة" means the head of this card's own دائرة. So each card is linked
to its parent in two steps: read the role, then find the card that heads that
unit. When the unit has no head card, a placeholder head is created and marked,
because a missing head card is itself a finding.
"""
import csv, re, sys, pathlib, collections, unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from parse_cards import cards, card_key

docx_root, map_csv, dst = map(pathlib.Path, sys.argv[1:4])
units = {m['card_key']: m['unit_bylaws'] for m in csv.DictReader(open(map_csv, encoding='utf-8')) if m['card_key']}


def key(s):
    """Unit name for matching: unify alef/taa/yaa, drop ال, unit words and spaces."""
    s = unicodedata.normalize('NFKC', s or '')
    s = re.sub('[أإآ]', 'ا', s).replace('ة', 'ه').replace('ى', 'ي')
    s = re.sub(r'\b(مديريه|دائره|الدائره|شعبه|الشعبه|قسم|القسم|مكتب|المكتب|فرع|الفرع)\b', ' ', s)
    s = re.sub(r'\bال', '', s)
    return re.sub(r'[\s\-–()/.]+', '', s)


# ---- 1. cards with their unit path -------------------------------------------
P = []
for f in sorted(docx_root.rglob('*.docx')):
    st = unicodedata.normalize('NFKC', f.stem)
    if 'بطاق' not in st and not st.startswith('2020'): continue
    art = card_key(f)
    for i, c in enumerate(cards(f), 1):
        if c['dept'] and re.match(r'(ال)?(قسم|فرع)\b', c['dept']) and not c['sub']:
            c['sub'], c['dept'] = c['dept'], None      # «قسم الشؤون الفنية» written as a دائرة
        P.append({'id': f'{art}/{f.stem[:12]}/{i}', 'art': art, **c,
                  'lv': {k: key(c[k]) for k in ('dept', 'sec', 'sub')}})

# readable unit names, so a placeholder can say which unit it heads
disp = {}
for p in P:
    for k in ('dept', 'sec', 'sub'):
        if p[k]: disp.setdefault((p['art'], p['lv'][k]), p[k])

# Order of the levels under a directorate. Usually دائرة > شعبة > قسم/مكتب.
# Two directorates put the قسم/فرع level ABOVE the دائرة, as their cards show:
#   17 الخدمات الفنية: المدير > قسم (معاون مدير ورئيس قسم) > دائرة > شعبة
#   38 فوج الإطفاء:    قائد الفوج > فرع > شعبة/ورشة/غرفة
ORDER = collections.defaultdict(lambda: ('dept', 'sec', 'sub'), {'17': ('sub', 'dept', 'sec'),
                                                                '38': ('sub', 'dept', 'sec')})

def unit(p, kind):
    """The unit at `kind` on this card's path: (article, level1, level2, ...)."""
    o = ORDER[p['art']]
    return (p['art'],) + tuple(p['lv'][k] for k in o[:o.index(kind) + 1])

# ---- 2. which unit, if any, each card heads -----------------------------------
DEP = r'معاون\s*(ال)?(مدير|قائد)'
SUB = r'(قسم|مكتب|فرع|ورشة|غرفة)'
def heads(p):
    t = p['title']; lv = p['lv']
    if re.match(DEP, t) and re.search(r'رئيس(ة)?\s*(ال)?' + SUB, t) and lv['sub']: return 'sub'
    if re.match(DEP, t): return 'deputy'
    if re.match(r'(ال)?مدير|قائد\s*(ال)?فوج|أمين السر|امين السر', t) and not any(lv.values()): return 'dir'
    for kind, w in (('sub', SUB), ('sec', 'شعبة'), ('dept', 'دائرة')):
        if re.match(r'رئيس(ة)?\s*(ال)?' + w, t) and lv[kind]: return kind
    return ''

head, head_kind = {}, {}
deputies = collections.defaultdict(list)
for p in P:
    p['heads'] = heads(p)
    if p['heads'] == 'deputy' or re.match(DEP, p['title']): deputies[p['art']].append(p)
    if p['heads'] in ('', 'deputy'): continue
    u = (p['art'],) if p['heads'] == 'dir' else unit(p, p['heads'])
    p['unit'] = u
    if u not in head: head[u], head_kind[u] = p['id'], p['heads']

# ---- 3. resolve the supervisor ------------------------------------------------
virtual = {}
def node_for(u, label):
    """Head card of a unit, or a placeholder head when no card describes it."""
    if u in head: return head[u], 'card'
    vid = 'VIRTUAL:' + '/'.join(x for x in u if x)
    name = disp.get((u[0], u[-1]), '') if len(u) > 1 else ''
    virtual.setdefault(vid, {'id': vid, 'art': u[0], 'unit': u,
                             'title': f'{label} — {name} (لا توجد بطاقة)' if name else f'{label} (لا توجد بطاقة)'})
    return vid, 'placeholder'

def by_name(kind, name, a):
    """A named unit in the same directorate: 'رئيس دائرة المشاتل' -> the head of المشاتل."""
    n = key(name)
    if len(n) < 3: return None
    hits = [pid for u, pid in head.items() if u[0] == a and head_kind[u] == kind and u[-1]
            and (u[-1] == n or n in u[-1] or u[-1] in n)]
    return hits[0] if len(hits) == 1 else None

def up(unit):
    u = list(unit)
    while len(u) > 2:
        u.pop()
        if u[-1]: return tuple(u)
    return (unit[0],) if len(unit) > 1 else None

def parent(p):
    b = (p['boss'] or '').strip(); a = p['art']
    if not b:
        # no supervisor written, but the card heads a unit: report to the unit above
        if p.get('unit'):
            above = up(p['unit'])
            if above: return node_for(above, 'رئيس الوحدة الأعلى')[0], 'مستنتج من الهيكل (لا رئيس في البطاقة)'
        return None, 'لا يوجد رئيس مباشر في البطاقة'
    if re.match(r'(السيد\s*)?المحافظ', b): return 'GOV', 'المحافظ'
    if re.search(r'أمين السر|امين السر', b): return node_for(('16.0',), 'أمين السر')[0], 'أمين السر'
    if re.match(DEP, b):
        ds = [x for x in deputies[a] if x['id'] != p['id']]
        if len(ds) == 1: return ds[0]['id'], 'معاون المدير'
        rest = key(re.sub(DEP, '', b))
        hit = [x for x in ds if len(rest) > 2 and (rest in key(x['title']) or rest in key(x['dept']))]
        if len(hit) == 1: return hit[0]['id'], 'معاون المدير (بالاسم)'
        return node_for((a,), 'المدير')[0], 'معاون المدير غير محدد ← المدير'
    if re.match(r'(ال)?مدير|قائد\s*(ال)?فوج', b): return node_for((a,), 'المدير')[0], 'المدير'
    m = re.match(r'(?:ر\.?\s*(?:ال)?ش\S*|رئيس(?:ة)?\s*(?:ال)?(دائرة|شعبة|' + SUB + r'))\s*(.*)', b)
    if m:
        kind = 'sec' if re.match(r'ر\.\s*(ال)?ش', b) else {'دائرة': 'dept', 'شعبة': 'sec'}.get(m.group(1), 'sub')
        named = by_name(kind, m.group(m.lastindex), a)
        if named and named != p['id']: return named, 'رئيس الوحدة بالاسم'
        if not p['lv'][kind]:
            # e.g. «رئيس الشعبة» on a card that names its دائرة but not the شعبة:
            # the nearest unit the card does name is the best we can say
            o = ORDER[a]; named_above = [k for k in o[:o.index(kind)] if p['lv'][k]]
            if named_above:
                return node_for(unit(p, named_above[-1]), 'رئيس الوحدة')[0], 'تقريبي: أقرب وحدة مذكورة'
            return None, 'الوحدة الأعلى غير مذكورة في البطاقة'
        u = unit(p, kind)
        if p.get('unit') == u: return None, 'يرأس الوحدة نفسها التي ذكرها كرئيس'
        pid, how = node_for(u, b)
        return pid, 'رئيس الوحدة من موقع البطاقة' + (' (بلا بطاقة)' if how == 'placeholder' else '')
    # anything else: a named position in the same directorate («قائد السرية», «امين مستودع»)
    hit = [x for x in P if x['art'] == a and x['id'] != p['id'] and key(x['title']) == key(b)]
    if len(hit) == 1: return hit[0]['id'], 'مسمى مطابق'
    return None, 'لم يُطابق'

for p in P:
    p['parent'], p['how'] = parent(p)

# Placeholders report to the head of the unit above them; a directorate head to
# the governor. Linking one placeholder can create another, so loop to a fixpoint.
while any('parent' not in v for v in virtual.values()):
    for v in [v for v in virtual.values() if 'parent' not in v]:
        above = up(v['unit'])
        v['parent'], v['how'] = (node_for(above, 'رئيس الوحدة الأعلى')[0], 'الوحدة الأعلى') if above \
            else ('GOV', 'رأس مديرية')

# ---- 4. report ----------------------------------------------------------------
how = collections.Counter('مرتبط' if p['parent'] else p['how'] for p in P)
via = collections.Counter(p['how'] for p in P if p['parent'])
cards_linked = sum(1 for p in P if p['parent'])
to_placeholder = sum(1 for p in P if p['parent'] and str(p['parent']).startswith('VIRTUAL'))
print(f'{len(P)} cards, {cards_linked} linked ({cards_linked * 100 // len(P)}%), '
      f'{to_placeholder} of them to a placeholder head; {len(virtual)} placeholder heads')
for k, v in how.most_common(): print(f'  {v:4} {k}')
print('linked via:'); [print(f'  {v:4} {k}') for k, v in via.most_common(12)]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'الهيكل'; ws.sheet_view.rightToLeft = True
ws.append(['المعرف', 'المادة', 'المديرية', 'الدائرة', 'الشعبة', 'قسم/مكتب/فرع', 'المسمى', 'يرأس',
           'الرئيس المباشر كما ورد', 'يتبع لـ', 'طريقة الربط'])
for c in ws[1]: c.font, c.fill = Font(bold=True), PatternFill('solid', fgColor='DDE7F0')
ws.append(['GOV', '', '', '', '', '', 'المحافظ', 'المحافظة', '', '', ''])
for v in virtual.values():
    u = v['unit'] + (None,) * (4 - len(v['unit']))
    ws.append([v['id'], v['art'], units.get(v['art']), u[1], u[2], u[3], v['title'], 'placeholder', '', v['parent'], v['how']])
for p in P:
    ws.append([p['id'], p['art'], units.get(p['art'], p['dir']), p['dept'], p['sec'], p['sub'], p['title'],
               p['heads'], p['boss'], p['parent'], p['how']])
for col, w in zip('ABCDEFGHIJK', [26, 7, 34, 24, 22, 18, 34, 10, 26, 26, 26]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wb.save(dst)
print('->', dst)

# ---- 5. the tree as a page ------------------------------------------------------
import html, json
nodes = [{'id': 'GOV', 'p': None, 't': 'المحافظ', 'u': 'محافظة دمشق', 'h': 'gov', 'f': ''}]
for v in virtual.values():
    nodes.append({'id': v['id'], 'p': v['parent'], 't': v['title'], 'u': '', 'h': 'placeholder', 'f': ''})
for p in P:
    loc = ' / '.join(x for x in (p['dept'], p['sec'], p['sub']) if x)
    flag = 'approx' if p['how'].startswith(('تقريبي', 'مستنتج', 'معاون المدير غير')) else ''
    nodes.append({'id': p['id'], 'p': p['parent'], 't': p['title'], 'u': loc, 'h': p['heads'], 'f': flag,
                  'a': p['art'], 'b': p['boss'] or '', 'w': p['how']})
for n in nodes:
    if n['id'] != 'GOV' and n['p'] is None: n['p'] = 'ORPHAN'
nodes.append({'id': 'ORPHAN', 'p': None, 't': 'بطاقات لم يُحدَّد رئيسها', 'u': '', 'h': 'orphan', 'f': ''})
for n in nodes:
    if n.get('a') and n['p'] == 'GOV': n['u'] = units.get(n['a'], '')
page = pathlib.Path(__file__).with_name('org_tree_template.html').read_text(encoding='utf-8')
out = pathlib.Path(dst).with_name(pathlib.Path(dst).stem.replace('hierarchy', 'org_tree') + '.html')
stats = {'cards': len(P), 'linked': cards_linked, 'placeholders': len(virtual),
         'approx': sum(1 for n in nodes if n['f'] == 'approx'),
         'orphans': sum(1 for n in nodes if n['p'] == 'ORPHAN')}
out.write_text(page.replace('/*DATA*/', json.dumps({'nodes': nodes, 'stats': stats}, ensure_ascii=False)),
               encoding='utf-8')
print('->', out)
